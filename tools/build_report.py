from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "project_report.md"
OUTPUT = ROOT / "project_report.docx"
ASSETS = ROOT / "report_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = "172B4D"
BLUE = "175CD3"
PALE = "EEF3FA"
PALE_ALT = "F7F9FC"
INK = "24364B"
MUTED = "52657A"
GREEN = "DDF3E7"
WHITE = "FFFFFF"
FONT_PATHS = [
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("C:/Windows/Fonts/segoeui.ttf"),
]


def font(size: int, bold: bool = False):
    candidates = ([Path("C:/Windows/Fonts/arialbd.ttf")] if bold else []) + FONT_PATHS
    for path in candidates:
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size)
            except OSError:
                pass
    return ImageFont.load_default()


def wrap(draw, text, fnt, max_width):
    lines = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        current = ""
        for word in words:
            candidate = word if not current else current + " " + word
            if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = word
        lines.append(current)
    return lines


def draw_centered(draw, box, text, fnt, fill, max_width=None, spacing=7):
    x1, y1, x2, y2 = box
    limit = max_width or (x2 - x1 - 36)
    lines = wrap(draw, text, fnt, limit)
    heights = [draw.textbbox((0, 0), line or " ", font=fnt)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + max(8, ((y2 - y1) - total) // 2)
    for line, h in zip(lines, heights):
        w = draw.textbbox((0, 0), line, font=fnt)[2]
        draw.text((x1 + ((x2 - x1) - w) / 2, y), line, font=fnt, fill=fill)
        y += h + spacing


def arrow(draw, start, end, color="#52657A", width=5, head=15):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    left = (end[0] - head * math.cos(angle - math.pi / 6), end[1] - head * math.sin(angle - math.pi / 6))
    right = (end[0] - head * math.cos(angle + math.pi / 6), end[1] - head * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=color)


def rounded_node(draw, box, text, fill, outline, fnt, text_fill="#172B4D", radius=22):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=3)
    draw_centered(draw, box, text, fnt, text_fill)


def make_context():
    image = Image.new("RGB", (1800, 1000), "#FFFFFF")
    d = ImageDraw.Draw(image)
    title_f = font(34, True)
    label_f = font(25, True)
    small_f = font(20)
    d.text((70, 38), "System context", fill=f"#{NAVY}", font=title_f)
    central = (610, 345, 1190, 655)
    rounded_node(d, central, "Internship and\nvacancy search\nweb application", f"#{BLUE}", f"#{BLUE}", font(34, True), "#FFFFFF", 36)
    nodes = {
        "seeker": (75, 130, 455, 315),
        "employer": (1345, 120, 1725, 300),
        "moderator": (75, 685, 455, 875),
        "source": (1345, 680, 1725, 870),
        "destination": (675, 800, 1125, 950),
    }
    rounded_node(d, nodes["seeker"], "Student / recent\ngraduate", "#F3F7FC", "#C6D4E3", label_f)
    rounded_node(d, nodes["employer"], "Employer", "#F3F7FC", "#C6D4E3", label_f)
    rounded_node(d, nodes["moderator"], "Moderator /\nadministrator", "#F3F7FC", "#C6D4E3", label_f)
    rounded_node(d, nodes["source"], "Authorized source or\nmanual curation", "#F3F7FC", "#C6D4E3", label_f)
    rounded_node(d, nodes["destination"], "Original application\ndestination", "#EEF8F3", "#ACD5BF", small_f)
    arrow(d, (455, 220), (610, 405))
    arrow(d, (610, 475), (455, 245), color="#175CD3", width=4)
    arrow(d, (1345, 215), (1190, 405))
    arrow(d, (455, 765), (610, 595))
    arrow(d, (1345, 765), (1190, 595))
    arrow(d, (900, 655), (900, 800))
    d.text((475, 286), "search / save", fill=f"#{MUTED}", font=small_f)
    d.text((465, 330), "results", fill=f"#{BLUE}", font=small_f)
    d.text((1210, 285), "submit listing", fill=f"#{MUTED}", font=small_f)
    d.text((455, 620), "review / publish", fill=f"#{MUTED}", font=small_f)
    d.text((1204, 618), "authorized records", fill=f"#{MUTED}", font=small_f)
    d.text((930, 708), "source link", fill=f"#{MUTED}", font=small_f)
    image.save(ASSETS / "context.png")


def draw_actor(draw, x, y, label):
    f = font(20, True)
    color = f"#{NAVY}"
    draw.ellipse((x - 23, y - 78, x + 23, y - 32), outline=color, width=4)
    draw.line((x, y - 32, x, y + 28), fill=color, width=4)
    draw.line((x - 34, y - 5, x + 34, y - 5), fill=color, width=4)
    draw.line((x, y + 28, x - 29, y + 73), fill=color, width=4)
    draw.line((x, y + 28, x + 29, y + 73), fill=color, width=4)
    w = draw.textbbox((0, 0), label, font=f)[2]
    draw.text((x - w / 2, y + 82), label, fill=color, font=f)


def make_use_case():
    image = Image.new("RGB", (1800, 1050), "#FFFFFF")
    d = ImageDraw.Draw(image)
    d.text((60, 38), "Use case diagram", fill=f"#{NAVY}", font=font(34, True))
    boundary = (380, 120, 1435, 940)
    d.rounded_rectangle(boundary, radius=28, outline="#AFC0D2", width=4, fill="#FBFCFE")
    d.text((420, 142), "Internship and vacancy search system", fill=f"#{BLUE}", font=font(25, True))
    actor_x_left, actor_x_right = 205, 1610
    draw_actor(d, actor_x_left, 295, "Job seeker")
    draw_actor(d, actor_x_right, 325, "Employer")
    draw_actor(d, actor_x_right, 730, "Moderator")
    usecases = [
        (535, 240, 950, 320, "Browse opportunities"),
        (1005, 240, 1370, 320, "Search and filter"),
        (535, 390, 950, 470, "View details and fit"),
        (1005, 390, 1370, 470, "Save opportunity"),
        (535, 540, 950, 620, "Open source link"),
        (1005, 540, 1370, 620, "Report incorrect listing"),
        (535, 710, 950, 790, "Submit opportunity"),
        (1005, 710, 1370, 790, "Approve / reject / expire"),
    ]
    for box in usecases:
        coords = box[:4]
        d.ellipse(coords, fill="#EEF4FF" if coords[1] < 680 else "#EAF5EF", outline="#A7BCE0" if coords[1] < 680 else "#A9CEB6", width=3)
        draw_centered(d, coords, box[4], font(20, True), f"#{INK}", spacing=5)
    # Actor-to-use-case links.
    for start, end in [((240, 295), (535, 280)), ((240, 310), (535, 430)), ((240, 325), (535, 580)), ((240, 320), (1005, 430)), ((240, 340), (1005, 580)), ((1575, 325), (1370, 750)), ((1575, 730), (1370, 750))]:
        arrow(d, start, end, color="#8292A6", width=3, head=11)
    image.save(ASSETS / "use_case.png")


def make_components():
    image = Image.new("RGB", (1800, 1050), "#FFFFFF")
    d = ImageDraw.Draw(image)
    d.text((60, 38), "Component architecture", fill=f"#{NAVY}", font=font(34, True))
    browser = (80, 410, 380, 600)
    api = (505, 390, 790, 620)
    db = (1430, 410, 1710, 600)
    rounded_node(d, browser, "Browser UI\nReact + TypeScript", "#EEF4FF", "#9DB8E7", font(24, True))
    rounded_node(d, api, "REST API\nFastAPI\nrole checks + validation", "#175CD3", "#175CD3", font(24, True), "#FFFFFF")
    rounded_node(d, db, "PostgreSQL\nusers, listings, skills,\nsaves, audit", "#EEF8F3", "#A9CEB6", font(22, True))
    boxes = [
        (920, 140, 1290, 255, "Identity and roles"),
        (920, 300, 1290, 415, "Catalogue and lifecycle"),
        (920, 460, 1290, 575, "Search, filters, fit rules"),
        (920, 620, 1290, 735, "Profile and saved items"),
        (920, 780, 1290, 895, "Employer + moderation"),
    ]
    for box in boxes:
        rounded_node(d, box[:4], box[4], "#F6F8FB", "#C6D4E3", font(21, True))
    # Client -> API and API -> modules.
    arrow(d, (380, 505), (505, 505), color=f"#{BLUE}", width=5)
    d.text((385, 465), "HTTPS / JSON", fill=f"#{MUTED}", font=font(17))
    for box in boxes:
        y = (box[1] + box[3]) // 2
        arrow(d, (790, 505), (920, y), color="#8292A6", width=3, head=11)
        arrow(d, (1290, y), (1430, 505), color="#8292A6", width=3, head=11)
    source = (80, 790, 380, 920)
    rounded_node(d, source, "Authorized import /\nemployer submission", "#FFF7E4", "#E6CE8D", font(21, True))
    arrow(d, (380, 855), (920, 835), color="#A3832D", width=3, head=12)
    audit = (505, 790, 790, 920)
    rounded_node(d, audit, "Moderation audit\nevents", "#F6F8FB", "#C6D4E3", font(21, True))
    arrow(d, (790, 850), (920, 835), color="#8292A6", width=3, head=11)
    image.save(ASSETS / "components.png")


def set_run_font(run, name="Arial", size=10, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, url):
    # Keep punctuation following a URL outside the hyperlink.
    visible = url
    tail = ""
    while visible and visible[-1] in ".,;:)":
        tail = visible[-1] + tail
        visible = visible[:-1]
    part = paragraph.part
    rid = part.relate_to(visible, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = visible
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    if tail:
        add_run(paragraph, tail)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|https?://[^\s]+)")


def add_run(paragraph, text, bold=False, italic=False, code=False):
    run = paragraph.add_run(text)
    set_run_font(run, name="Consolas" if code else "Arial", size=9.2 if code else 10.2, color=BLUE if code else INK, bold=bold, italic=italic)
    if code:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F1F4F8")
        run._element.get_or_add_rPr().append(shading)
    return run


def add_inline(paragraph, text):
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            add_run(paragraph, text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            add_run(paragraph, token[2:-2], bold=True)
        elif token.startswith("*"):
            add_run(paragraph, token[1:-1], italic=True)
        elif token.startswith("`"):
            add_run(paragraph, token[1:-1], code=True)
        else:
            add_hyperlink(paragraph, token)
        pos = match.end()
    if pos < len(text):
        add_run(paragraph, text[pos:])


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=95, bottom=90, end=95):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D7E0EA")


def parse_md_row(line):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def add_markdown_table(doc, rows):
    if len(rows) < 2:
        return
    header = parse_md_row(rows[0])
    body = [parse_md_row(r) for r in rows[1:] if not re.fullmatch(r"\|?\s*[-:| ]+\s*\|?", r)]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, text)
        for run in p.runs:
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(8.0)
            run.bold = True
    for ri, row in enumerate(body):
        cells = table.add_row().cells
        for ci in range(len(header)):
            cell = cells[ci]
            set_cell_margins(cell, top=75, start=85, bottom=75, end=85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ri % 2 == 1:
                set_cell_shading(cell, PALE_ALT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            add_inline(p, row[ci] if ci < len(row) else "")
            for run in p.runs:
                run.font.size = Pt(7.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.32)
    sec.different_first_page_header_footer = True
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 25, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11, NAVY)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(13 if name != "Title" else 30)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.space_after = Pt(14)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("AITU  ·  SE-2427  ·  Материалы аппробационного периода")
    set_run_font(r, size=8, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Проект дипломной работы  ·  25 октября 2026  ·  Стр. ")
    set_run_font(r, size=8, color=MUTED)
    add_page_number(footer)


def add_image(doc, target, alt):
    mapping = {
        "DIAGRAM_CONTEXT": ASSETS / "context.png",
        "DIAGRAM_USE_CASE": ASSETS / "use_case.png",
        "DIAGRAM_COMPONENT": ASSETS / "components.png",
    }
    image_path = mapping.get(target)
    if not image_path or not image_path.exists():
        p = doc.add_paragraph()
        add_run(p, f"Диаграмма не найдена: {target}", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(image_path), width=Inches(6.65))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    r = caption.add_run(alt)
    set_run_font(r, size=8.5, color=MUTED, italic=True)


def markdown_to_docx():
    make_context()
    make_use_case()
    make_components()
    doc = Document()
    configure_document(doc)
    lines = REPORT.read_text(encoding="utf-8").splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line == "---":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, rows)
            continue
        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image_match:
            add_image(doc, image_match.group(2), image_match.group(1))
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title" if first_heading else "Heading 1")
            add_inline(p, line[2:].strip())
            first_heading = False
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:].strip())
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:].strip())
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:].strip())
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line[2:])
            i += 1
            continue
        number_match = re.match(r"\d+\.\s+(.*)", line)
        if number_match:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, number_match.group(1))
            i += 1
            continue
        p = doc.add_paragraph()
        if line.startswith("**") and line.endswith("**") and line.count("**") == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
        add_inline(p, line.replace("  ", " "))
        i += 1
    doc.core_properties.title = "Разработка веб-приложения для поиска стажировок и вакансий"
    doc.core_properties.subject = "Материалы к аппробационному периоду"
    doc.core_properties.author = "Жиенбек Муслим Рахымжанұлы; Сағым Мұхаммеджан Бақытжанұлы"
    doc.core_properties.keywords = "job search, internship, diploma project, SE-2427"
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    markdown_to_docx()
